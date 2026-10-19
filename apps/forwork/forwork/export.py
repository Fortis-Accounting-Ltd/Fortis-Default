"""
Getting data out - a board as an Excel workbook (Board > Export to Excel), and everything as one readable
JSON file that goes into each nightly backup, so the firm's work can be read without ForWork.
"""
from __future__ import annotations

import gzip
import json
from pathlib import Path

from .db import now_iso
from .model import column_type, fmt_number, date_text


def cell_text(col: dict, item: dict, ctx, users: dict) -> str:
    t = col["type"]
    if t == "name":
        return item["name"]
    comp = (item.get("computed") or {}).get(col["id"])
    if t == "mirror":
        return ", ".join(c["text"] for c in comp or [])
    if t == "formula":
        return "" if comp is None or isinstance(comp, dict) else fmt_number(comp)
    if t == "subtasks":
        return f"{comp['done']}/{comp['total']}" if comp else ""
    if t == "board_relation":
        return ", ".join(x["name"] for x in comp or [])
    if t == "creation_log":
        return f"{users.get(item.get('created_by'), '')} {date_text((item.get('created_at') or '')[:10])}".strip()
    if t == "last_updated":
        return f"{users.get(item.get('updated_by'), '')} {date_text((item.get('updated_at') or '')[:10])}".strip()
    v = (item.get("values") or {}).get(col["id"])
    return column_type(t).text(v, col["settings"], ctx) if v is not None else ""


def board_rows(svc, bid: int) -> dict:
    p = svc.board_payload(bid)
    ctx = svc.ctx()
    users = ctx.users
    groups = {g["id"]: g for g in p["groups"]}
    item_cols = [c for c in p["columns"] if c["scope"] == "item" and c["type"] != "subtasks"]
    sub_cols = [c for c in p["columns"] if c["scope"] == "subitem"]
    items, subs = [], []
    for it in p["items"]:
        items.append([groups.get(it["group_id"], {}).get("title", "")] + [cell_text(c, it, ctx, users) for c in item_cols])
        for s in it.get("subitems") or []:
            subs.append([it["name"]] + [cell_text(c, s, ctx, users) for c in sub_cols])
    return {"board": p["board"], "item_headers": ["Group"] + [c["title"] for c in item_cols], "items": items,
            "sub_headers": ["Parent item"] + [c["title"] for c in sub_cols], "subitems": subs}


def board_xlsx(svc, bid: int, path: str | Path) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    data = board_rows(svc, bid)
    wb = Workbook()
    ws = wb.active
    ws.title = "Items"
    head_font, fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="0050C8")
    for sheet, headers, rows in ((ws, data["item_headers"], data["items"]),):
        sheet.append(headers)
        for cell in sheet[1]:
            cell.font, cell.fill = head_font, fill
        for r in rows:
            sheet.append(r)
        sheet.freeze_panes = "C2"
    if data["subitems"]:
        s2 = wb.create_sheet("Subitems")
        s2.append(data["sub_headers"])
        for cell in s2[1]:
            cell.font, cell.fill = head_font, fill
        for r in data["subitems"]:
            s2.append(r)
    for sheet in wb.worksheets:
        for col in sheet.columns:
            width = max((len(str(c.value or "")) for c in list(col)[:200]), default=8)
            sheet.column_dimensions[col[0].column_letter].width = min(60, max(10, width + 2))
    info = wb.create_sheet("About")
    info.append(["Board", data["board"]["name"]])
    info.append(["Exported", now_iso()])
    info.append(["From", "ForWork - Fortis Accounting Partners"])
    path = Path(path)
    wb.save(path)
    return path


def export_all_json(svc, dest_dir: str | Path) -> list[Path]:
    """Every board (active and archived) with its items, values as shown, and updates as plain text."""
    dest = Path(dest_dir)
    dest.mkdir(parents=True, exist_ok=True)
    stamp = now_iso().replace(":", "").replace("-", "")[:15]
    path = dest / f"forwork-export-{stamp}Z.json.gz"
    boards = []
    for b in svc.boards(states=("active", "archived")):
        d = board_rows(svc, b["id"])
        updates = []
        for r in svc.db.q("SELECT u.item_id, i.name AS item, u.parent_id, u.created_at, COALESCE(us.name, u.author_name) AS author, "
                          "u.text FROM updates u JOIN items i ON i.id = u.item_id LEFT JOIN users us ON us.id = u.created_by "
                          "WHERE i.board_id = ? AND u.state = 'active' ORDER BY u.created_at", (b["id"],)):
            updates.append(dict(r))
        boards.append({"id": b["id"], "name": b["name"], "state": b["state"],
                       "columns": d["item_headers"], "items": d["items"],
                       "subitem_columns": d["sub_headers"], "subitems": d["subitems"], "updates": updates})
    with gzip.open(path, "wt", encoding="utf-8") as f:
        json.dump({"exported_at": now_iso(), "app": "ForWork", "boards": boards}, f, ensure_ascii=False)
    return [path]
