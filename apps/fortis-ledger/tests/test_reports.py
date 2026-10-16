import datetime as dt

import openpyxl
from pypdf import PdfReader

from fortis_ledger import reports
from fortis_ledger.ledger import Ledger

from conftest import J


def pdf_text(path) -> str:
    return "\n".join(p.extract_text() or "" for p in PdfReader(str(path)).pages)


def two_years(L):
    L.close_year(2025)
    L.save_journal(J("2026-06-30", "FY2026 sales", ("13600", 60000, 0), ("51100", 0, 60000)), 2026)
    L.save_journal(J("2026-03-31", "Accrued fees", ("72930", 1500, 0), ("31300", 0, 1500)), 2026)
    return Ledger(L.store, L.slug)


def test_financial_reports_pack(tmp_path, company):
    L = two_years(company)
    res = reports.financial_reports(L, 2026, tmp_path / "fr.pdf", final=True, date_signed="15 October 2026")
    assert not res["has_errors"] and res["draft"] is False
    assert res["summary"]["npat"]["cy"] == 58500 and res["summary"]["npat"]["py"] == 35000
    text = pdf_text(tmp_path / "fr.pdf")
    assert "Test Co Pty Ltd" in text and "Balance Sheet" in text and "58,500" in text
    assert "DRAFT" not in text
    draft = reports.financial_reports(L, 2026, tmp_path / "draft.pdf")
    assert draft["draft"] is True and "DRAFT" in pdf_text(tmp_path / "draft.pdf")
    acts = [a for a in L.store.read_audit(L.slug) if a["action"] == "report.generated"]
    assert {a["detail"]["file"] for a in acts} == {"fr.pdf", "draft.pdf"}


def test_trial_balance_pdf_and_excel(tmp_path, company):
    L = two_years(company)
    reports.trial_balance(L, 2026, tmp_path / "tb.pdf")
    text = pdf_text(tmp_path / "tb.pdf")
    assert "Trial Balance" in text and "As at 30 June 2026" in text and "Trade debtors" in text
    reports.trial_balance(L, 2026, tmp_path / "tb.xlsx", as_excel=True)
    ws = openpyxl.load_workbook(tmp_path / "tb.xlsx").active
    rows = [r for r in ws.iter_rows(min_row=6, values_only=True)]
    total = rows[-1]
    assert total[1] == "Total" and total[4] == total[5] and total[4] > 0
    debtors = next(r for r in rows if r[0] == "13600")
    assert debtors[4] == 105000


def test_general_ledger_running_balance(tmp_path, company):
    L = two_years(company)
    gl = L.general_ledger(2026, codes=["31300"])
    assert len(gl) == 1 and gl[0]["opening"] == -4900 and gl[0]["closing"] == -6400
    reports.general_ledger(L, 2026, tmp_path / "gl.xlsx", codes=["31300", "13600"], as_excel=True)
    ws = openpyxl.load_workbook(tmp_path / "gl.xlsx").active
    vals = [r for r in ws.iter_rows(min_row=6, values_only=True)]
    assert any(r[2] and str(r[2]).startswith("13600") for r in vals)
    reports.general_ledger(L, 2026, tmp_path / "gl.pdf", start=dt.date(2026, 4, 1))
    text = pdf_text(tmp_path / "gl.pdf")
    assert "1 April 2026 to 30 June 2026" in text and "Accrued fees" not in text   # before the period: in the opening


def test_journal_report_shows_void(tmp_path, company):
    L = two_years(company)
    j = L.save_journal(J("2026-06-30", "Wrong entry", ("72930", 10, 0), ("31300", 0, 10)), 2026)
    L.void_journal(2026, j.number, "keyed twice")
    reports.journal_report(L, 2026, tmp_path / "jr.pdf")
    text = pdf_text(tmp_path / "jr.pdf")
    assert "Wrong entry (VOID)" in text and "FY2026 sales" in text and "Accrued fees" in text
    reports.journal_report(L, 2026, tmp_path / "jr2.pdf", include_void=False)
    assert "Wrong entry" not in pdf_text(tmp_path / "jr2.pdf")
