"""Profit and Loss, Balance Sheet and the GST report: sections, comparatives, cash and accrual GST, the files."""
import datetime as dt
from decimal import Decimal

import openpyxl
import pytest
from pypdf import PdfReader

from fortis_ledger import bank, reports, statements
from fortis_ledger.ledger import Ledger
from fortis_ledger.model import Journal, Line
from fortis_ledger.money import to_dec

from conftest import J

D = Decimal


def rows_by_label(data) -> dict:
    return {r["label"]: r for r in data["rows"]}


def pdf_text(path) -> str:
    return "\n".join(p.extract_text() or "" for p in PdfReader(str(path)).pages)


@pytest.fixture
def two_years(company):
    """FY2025 from the conftest balances, rolled forward, and some FY2026 trading."""
    L = company
    L.close_year(2025)
    L = Ledger(L.store, L.slug)
    cos = next(a.code for a in L.chart.usable() if a.report_code.startswith("EXP.COS"))
    other = next(a.code for a in L.chart.usable() if a.report_code.startswith("REV.INV.INT"))
    L.save_journal(J("2025-09-30", "Q1 sales", ("13600", 30000, 0), ("51100", 0, 30000)), 2026)
    L.save_journal(J("2026-03-31", "Q3 sales", ("13600", 30000, 0), ("51100", 0, 30000)), 2026)
    L.save_journal(J("2026-03-31", "Stock bought", (cos, 12000, 0), ("31300", 0, 12000)), 2026)
    L.save_journal(J("2026-06-30", "Interest", ("18500", 250, 0), (other, 0, 250)), 2026)
    L.save_journal(J("2026-06-30", "Fees", ("72930", 1500, 0), ("31300", 0, 1500)), 2026)
    L.save_journal(J("2026-06-30", "Tax", ("99100", 5000, 0), ("35300", 0, 5000)), 2026)
    L = Ledger(L.store, L.slug)
    L.cos, L.other = cos, other
    return L


def test_profit_and_loss_sections_and_comparatives(two_years):
    L = two_years
    data = statements.profit_and_loss(L, *L.fy_range(2026))
    assert data["columns"] == ["FY2026", "FY2025"] and data["period"] == "For the year ended 30 June 2026"
    r = rows_by_label(data)
    assert [x["label"] for x in data["rows"] if x["kind"] == "heading"] == [
        "Trading Income", "Cost of Sales", "Other Income", "Operating Expenses", "Income Tax Expense"]
    assert r["Total Trading Income"]["cy"] == D("60000.00") and r["Total Trading Income"]["py"] == D("50000.00")
    assert r["Gross Profit"]["cy"] == D("48000.00")
    assert r["Total Other Income"]["cy"] == D("250.00")
    assert r["Net Profit Before Income Tax"]["cy"] == D("46750.00")
    assert r["Net Profit After Income Tax"]["cy"] == D("41750.00") and r["Net Profit After Income Tax"]["py"] == D("35000.00")
    assert data["net_profit"] == L.trial_balance(2026)["profit"]
    # part of a year, compared with the same months a year earlier
    half = statements.profit_and_loss(L, dt.date(2025, 7, 1), dt.date(2025, 12, 31))
    assert half["columns"] == ["31 Dec 2025", "31 Dec 2024"] and half["period"].startswith("For the period 1 July 2025")
    assert rows_by_label(half)["Total Trading Income"]["cy"] == D("30000.00")
    assert rows_by_label(half)["Total Trading Income"]["py"] == D("0")
    only = statements.profit_and_loss(L, *L.fy_range(2026), compare=False)
    assert only["columns"] == ["FY2026"] and all(x["py"] is None for x in only["rows"])


def test_profit_and_loss_across_years_and_distributions(store):
    c = store.create_client("Family Trust", "trust")
    L = Ledger(store, c.slug)
    for fy in (2025, 2026):
        L.ensure_year(fy)
    dist = next(a.code for a in L.chart.usable() if a.report_code.startswith("EXP.DIS"))
    sales = next(a.code for a in L.chart.usable() if a.report_code.startswith("REV.TRA"))
    bank_code = next(a.code for a in L.chart.usable() if a.report_code.startswith("ASS.CUR.CAS.BAN"))
    L.save_journal(J("2025-05-31", "May", (bank_code, 1000, 0), (sales, 0, 1000)), 2025)
    L.save_journal(J("2025-07-31", "July", (bank_code, 2500, 0), (sales, 0, 2500)), 2026)
    L.save_journal(J("2025-06-30", "Distribution", (dist, 1000, 0), (bank_code, 0, 1000)), 2025)
    L = Ledger(store, c.slug)
    data = statements.profit_and_loss(L, dt.date(2025, 5, 1), dt.date(2025, 7, 31), compare=False)
    r = rows_by_label(data)
    assert r["Total Trading Income"]["cy"] == D("3500.00")          # May (FY2025) and July (FY2026)
    assert r["Total Distributions"]["cy"] == D("1000.00")
    assert r["Net Profit"]["cy"] == D("3500.00") and r["Undistributed Profit"]["cy"] == D("2500.00")
    with pytest.raises(statements.StatementError, match="not been started"):
        statements.profit_and_loss(L, dt.date(2026, 5, 1), dt.date(2026, 7, 31))


def test_balance_sheet_ties_to_the_trial_balance(two_years):
    L = two_years
    data = statements.balance_sheet(L, L.fy_range(2026)[1])
    assert data["columns"] == ["30 Jun 2026", "30 Jun 2025"] and data["out_by"] == 0
    r = rows_by_label(data)
    tb = {row["code"]: row["balance"] for row in L.trial_balance(2026)["rows"]}
    assert r["Total Bank"]["cy"] == tb["18500"] == D("15250.00")
    assert r["Total Current Assets"]["cy"] == D("105000.00")        # debtors 45,000 + 60,000
    assert r["Total Assets"]["cy"] == D("120250.00")
    assert r["Total Current Liabilities"]["cy"] == D("33400.00")     # creditors 4,900 + 13,500, tax 15,000
    assert r["Net Assets"]["cy"] == r["Total Equity"]["cy"] == D("86850.00")
    assert r["Current Year Earnings"]["cy"] == D("41750.00") and r["Current Year Earnings"]["py"] == D("35000.00")
    assert r["Retained Earnings"]["cy"] == D("45000.00")             # 10,000 + FY2025's 35,000
    assert r["Net Assets"]["py"] == D("45100.00")
    mid = statements.balance_sheet(L, dt.date(2025, 12, 31))
    assert mid["columns"] == ["31 Dec 2025", "31 Dec 2024"] and mid["compare"]
    assert rows_by_label(mid)["Current Year Earnings"]["cy"] == D("30000.00")
    assert rows_by_label(mid)["Net Assets"]["py"] == 0                 # FY2025's balances are dated 30 June 2025
    with pytest.raises(statements.StatementError, match="FY2027"):
        statements.balance_sheet(L, dt.date(2026, 7, 1))


def gst_ledger(store):
    c = store.create_client("GST Co Pty Ltd", "company", gst_registered=True)
    L = Ledger(store, c.slug)
    L.ensure_year(2026)
    L = Ledger(store, c.slug)
    feed_lines = [
        {"date": dt.date(2026, 4, 2), "amount": D("2200.00"), "payee": "", "description": "ACME PTY LTD INV 1001", "reference": "", "balance": None},
        {"date": dt.date(2026, 4, 3), "amount": D("-110.00"), "payee": "", "description": "OFFICEWORKS", "reference": "", "balance": None},
        {"date": dt.date(2026, 4, 4), "amount": D("-3300.00"), "payee": "", "description": "DELL LAPTOP", "reference": "", "balance": None},
        {"date": dt.date(2026, 4, 5), "amount": D("-22.00"), "payee": "", "description": "OFFICEWORKS REFUND", "reference": "", "balance": None},
        {"date": dt.date(2026, 4, 6), "amount": D("500.00"), "payee": "", "description": "RENT RECEIVED", "reference": "", "balance": None},
        {"date": dt.date(2026, 4, 7), "amount": D("-9.00"), "payee": "", "description": "BANK FEE", "reference": "", "balance": None},
        {"date": dt.date(2026, 7, 1), "amount": D("-55.00"), "payee": "", "description": "NEXT QUARTER", "reference": "", "balance": None},
    ]
    feed_lines[3]["amount"] = D("22.00")                  # money back from the supplier
    bank.import_lines(L, "18500", feed_lines, "stmt.csv", bank.Layout())
    ids = [ln["id"] for ln in bank.BankFeed(store, c.slug, "18500").ordered()]
    fixed = next(a.code for a in L.chart.usable() if a.report_code.startswith("ASS.NCA.FIX"))
    res = bank.reconcile(L, "18500", [bank.Coding(ids[0], "51100", "GST on Income"),
                                      bank.Coding(ids[1], "72100", "GST on Expenses"),
                                      bank.Coding(ids[2], fixed, "GST on Capital"),
                                      bank.Coding(ids[3], "72100", "GST on Expenses"),
                                      bank.Coding(ids[4], "51100", "Input Taxed"),
                                      bank.Coding(ids[5], "72500", "BAS Excluded")])
    assert not res["errors"], res["errors"]
    # an invoice entered as a journal, GST inclusive: accrual basis only
    j = Journal(date=dt.date(2026, 5, 15), narration="Invoice 1002 to Beta Pty Ltd",
                lines=[Line("13600", debit=to_dec(1100)), Line("51100", credit=to_dec(1100), tax="GST on Income")])
    L.apply_gst(j)
    L.save_journal(j, 2026)
    return Ledger(store, c.slug)


def test_gst_report_cash_and_accrual(store):
    L = gst_ledger(store)
    start, end = dt.date(2026, 4, 1), dt.date(2026, 6, 30)
    acc = statements.gst_report(L, start, end, "accrual")
    labels = {k: v for k, _, v in acc["labels"]}
    assert labels["G1"] == D("3800.00")                  # 2,200 + 500 input taxed rent + 1,100 invoice
    assert labels["G4"] == D("500.00") and labels["G3"] == 0
    assert labels["G10"] == D("3300.00")
    assert labels["G11"] == D("88.00")                   # 110 less the 22 refund; the BAS excluded fee is not in it
    assert labels["1A"] == D("300.00")                   # 200 + 100 on the invoice
    assert labels["1B"] == D("308.00")                   # 10 + 300 - 2 refunded
    assert acc["net_gst"] == D("-8.00") and acc["left_out"] == 0
    cash = statements.gst_report(L, start, end, "cash")
    cl = {k: v for k, _, v in cash["labels"]}
    assert cl["G1"] == D("2700.00") and cl["1A"] == D("200.00") and cash["left_out"] == 1
    assert all(x["kind"] == "bank" for x in cash["lines"])
    refund = next(x for x in acc["lines"] if x["text"].startswith("OFFICEWORKS REFUND"))
    assert (refund["net"], refund["gst"], refund["gross"]) == (D("-20.00"), D("-2.00"), D("-22.00"))
    gst_code = bank.gst_account(L)
    assert L.closing(2026)[gst_code] == D("8.00")      # the GST account agrees: 308 paid less 300 collected
    assert not [x for x in acc["lines"] if x["date"] > end]
    with pytest.raises(statements.StatementError):
        statements.gst_report(L, start, end, "monthly")


def test_statement_and_gst_files(tmp_path, two_years, store):
    L = two_years
    reports.profit_and_loss(L, tmp_path / "pl.pdf", *L.fy_range(2026))
    text = pdf_text(tmp_path / "pl.pdf")
    assert "Profit and Loss" in text and "Gross Profit" in text and "41,750.00" in text and "FY2025" in text
    reports.balance_sheet(L, tmp_path / "bs.xlsx", L.fy_range(2026)[1], as_excel=True)
    ws = openpyxl.load_workbook(tmp_path / "bs.xlsx").active
    cells = {r[0]: r[1:] for r in ws.iter_rows(values_only=True) if r and r[0]}
    assert cells["Net Assets"][0] == 86850 and cells["Total Equity"][0] == 86850
    assert ws["A1"].value == "Balance Sheet" and cells["Net Assets"][1] == 45100
    G = gst_ledger(store)
    reports.gst_report(G, tmp_path / "gst.pdf", dt.date(2026, 4, 1), dt.date(2026, 6, 30), "cash")
    text = pdf_text(tmp_path / "gst.pdf")
    assert "GST Report" in text and "Cash basis" in text and "GST refundable" in text and "left out" in text
    reports.gst_report(G, tmp_path / "gst.xlsx", dt.date(2026, 4, 1), dt.date(2026, 6, 30), "accrual", as_excel=True)
    wb = openpyxl.load_workbook(tmp_path / "gst.xlsx")
    assert wb.sheetnames == ["BAS summary", "By tax rate", "Transactions"]
    bas = {r[0]: r[2] for r in wb["BAS summary"].iter_rows(min_row=6, values_only=True) if r and r[0]}
    assert bas["1A"] == 300 and bas["1B"] == 308
    actions = [a["detail"].get("report") for a in G.store.read_audit(G.slug) if a["action"] == "report.generated"]
    assert actions[:2] == ["gst_report", "gst_report"]
