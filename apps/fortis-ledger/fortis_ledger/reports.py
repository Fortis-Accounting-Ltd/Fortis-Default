"""
Reports.

* Financial Reports pack - the firm's Xero report template look, produced by the report
  engine (fortis_ledger/reportengine) from the ledger's balances (current and prior year).
* Trial Balance, Profit and Loss, Balance Sheet, General Ledger Detail (account transactions), GST
  Report (cash or accrual basis) and Journal Report - Xero-style tables, as PDF or Excel.
"""
from __future__ import annotations

import copy
import datetime as dt
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.pdfgen import canvas as rl_canvas
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from . import statements as stm
from .engine import BOILERPLATE, REPORT_MAP, fonts, fr_engine, render_pack
from .ledger import Ledger
from .money import ZERO, fmt

BLUE = colors.HexColor("#06B3E8")
RULE = colors.HexColor("#D9D9D9")
DARK = colors.HexColor("#8C8C8C")
NUM_FMT = '#,##0.00;(#,##0.00);"-"'


# ---------------------------------------------------------------------------------------- #
# Financial Reports pack
# ---------------------------------------------------------------------------------------- #
def financial_reports(ledger: Ledger, fy: int, out_path: str | Path, final: bool = False,
                      upto: dt.date | None = None, date_signed: str = "", order: list[str] | None = None) -> dict:
    data = ledger.pack(fy, upto, date_signed)
    pack = fr_engine.Pack(data, REPORT_MAP)
    boiler = copy.deepcopy(BOILERPLATE)
    for k, v in (ledger.client.options.get("boilerplate_overrides") or {}).items():
        if isinstance(v, dict) and isinstance(boiler.get(k), dict):
            boiler[k].update(v)
        else:
            boiler[k] = v
    draft = (not final) or pack.has_errors()
    render_pack.build_pdf(pack, out_path, boiler, fonts(), order, draft=draft)
    ledger.store.audit(ledger.slug, "report.generated", {"fy": fy, "report": "financial_reports", "final": not draft,
                                                         "file": Path(out_path).name})
    return {"path": str(out_path), "draft": draft, "validation": pack.validation, "summary": pack.summary(),
            "has_errors": pack.has_errors()}


# ---------------------------------------------------------------------------------------- #
# table reports - common PDF / Excel builders
# ---------------------------------------------------------------------------------------- #
class _Canvas(rl_canvas.Canvas):
    """Footer 'Report | Entity' and 'Page x of y' on every page."""

    def __init__(self, *a, footer_left: str = "", font: str = "Helvetica", **k):
        super().__init__(*a, **k)
        self._pages = []
        self._footer_left = footer_left
        self._font = font

    def showPage(self):
        self._pages.append(dict(self.__dict__))
        self._startPage()

    def save(self):
        total = len(self._pages)
        for st in self._pages:
            self.__dict__.update(st)
            w = self._pagesize[0]
            self.setFont(self._font, 7)
            self.setFillColor(colors.black)
            self.drawString(17 * mm, 12 * mm, self._footer_left)
            self.drawRightString(w - 17 * mm, 12 * mm, f"Page {self._pageNumber} of {total}")
            self.setStrokeColor(RULE)
            self.line(17 * mm, 16 * mm, w - 17 * mm, 16 * mm)
            rl_canvas.Canvas.showPage(self)
        rl_canvas.Canvas.save(self)


def _styles():
    f = fonts()
    return {
        "title": ParagraphStyle("t", fontName=f["SB"], fontSize=20, leading=24, textColor=BLUE, spaceAfter=2),
        "sub": ParagraphStyle("s", fontName=f["SB"], fontSize=12, leading=15),
        "cell": ParagraphStyle("c", fontName=f["R"], fontSize=7.5, leading=9),
        "cellb": ParagraphStyle("cb", fontName=f["SB"], fontSize=7.5, leading=9),
        "head": ParagraphStyle("h", fontName=f["SB"], fontSize=7, leading=9),
    }


BOLD_KINDS = ("total", "heading", "subheading", "result", "grand")


def _pdf_table(out_path, title: str, entity: str, period: str, header: list[str], rows: list[dict],
               widths: list[float], num_cols: set[int], landscape_page: bool = False) -> str:
    """rows: {'cells': [...], 'kind': 'line' | 'total' | 'heading' | 'subheading' | 'result' | 'grand'}"""
    return _pdf_tables(out_path, title, entity, period,
                       [{"header": header, "rows": rows, "widths": widths, "num_cols": num_cols}], landscape_page)


def _pdf_tables(out_path, title: str, entity: str, period: str, tables: list[dict], landscape_page: bool = False,
                notes: list[str] | None = None) -> str:
    """One or more tables under the report title: {'heading': ..., 'header', 'rows', 'widths', 'num_cols'}."""
    st = _styles()
    size = landscape(A4) if landscape_page else A4
    doc = SimpleDocTemplate(str(out_path), pagesize=size, leftMargin=17 * mm, rightMargin=17 * mm,
                            topMargin=18 * mm, bottomMargin=22 * mm, title=f"{title} - {entity}")
    story = [Paragraph(title, st["title"]), Paragraph(_esc(entity), st["sub"]), Paragraph(_esc(period), st["sub"]),
             Spacer(1, 6 * mm)]
    for tab in tables:
        if tab.get("heading"):
            story += [Spacer(1, 3 * mm), Paragraph(_esc(tab["heading"]), st["sub"]), Spacer(1, 2 * mm)]
        story.append(_pdf_grid(st, tab["header"], tab["rows"], tab["widths"], tab["num_cols"]))
    for n in notes or []:
        story += [Spacer(1, 3 * mm), Paragraph(_esc(n), st["cell"])]
    f = fonts()
    doc.build(story, canvasmaker=lambda *a, **k: _Canvas(*a, footer_left=f"{title}  |  {entity}", font=f["R"], **k))
    return str(out_path)


def _pdf_grid(st, header: list[str], rows: list[dict], widths: list[float], num_cols: set[int]) -> Table:
    data = [[Paragraph(h.upper(), ParagraphStyle("hn", parent=st["head"], alignment=2 if ci in num_cols else 0))
             for ci, h in enumerate(header)]]
    styles = [("LINEBELOW", (0, 0), (-1, 0), 0.8, DARK), ("VALIGN", (0, 0), (-1, -1), "TOP"),
              ("TOPPADDING", (0, 0), (-1, -1), 2), ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
              ("LEFTPADDING", (0, 0), (-1, -1), 2), ("RIGHTPADDING", (0, 0), (-1, -1), 2)]
    for ci in num_cols:
        styles.append(("ALIGN", (ci, 0), (ci, -1), "RIGHT"))
    for r in rows:
        i = len(data)
        kind = r.get("kind", "line")
        sty = st["cellb"] if kind in BOLD_KINDS else st["cell"]
        cells = []
        for ci, v in enumerate(r["cells"]):
            if ci in num_cols:
                cells.append(Paragraph(v, ParagraphStyle("n", parent=sty, alignment=2)))
            else:
                cells.append(Paragraph(_esc(v), ParagraphStyle("i", parent=sty, leftIndent=r.get("indent", 0) * mm)))
        data.append(cells)
        if kind == "line":
            styles.append(("LINEBELOW", (0, i), (-1, i), 0.3, RULE))
        elif kind in ("total", "result"):
            styles.append(("LINEABOVE", (0, i), (-1, i), 0.6, DARK))
            styles.append(("LINEBELOW", (0, i), (-1, i), 0.6, DARK))
        elif kind == "grand":
            styles.append(("LINEABOVE", (0, i), (-1, i), 0.6, DARK))
            styles.append(("LINEBELOW", (0, i), (-1, i), 1.6, DARK))
        elif kind == "heading":
            styles.append(("TOPPADDING", (0, i), (-1, i), 8))
        elif kind == "subheading":
            styles.append(("TOPPADDING", (0, i), (-1, i), 5))
    t = Table(data, colWidths=widths, repeatRows=1)
    t.setStyle(TableStyle(styles))
    return t


def _xlsx(out_path, title: str, entity: str, period: str, header: list[str], rows: list[dict],
          num_cols: set[int], widths: list[int]) -> str:
    return _xlsx_sheets(out_path, title, entity, period,
                        [{"title": title, "header": header, "rows": rows, "num_cols": num_cols, "widths": widths}])


def _xlsx_sheets(out_path, title: str, entity: str, period: str, sheets: list[dict], notes: list[str] | None = None) -> str:
    wb = Workbook()
    for k, sh in enumerate(sheets):
        ws = wb.active if k == 0 else wb.create_sheet()
        ws.title = sh["title"][:31]
        ws.append([title if k == 0 else sh["title"]])
        ws.append([entity])
        ws.append([period])
        ws.append([])
        ws.append(sh["header"])
        ws["A1"].font = Font(bold=True, size=14, color="06B3E8")
        for c in ws[5]:
            c.font = Font(bold=True)
            c.border = Border(bottom=Side(style="thin"))
        for r in sh["rows"]:
            ws.append([float(v) if isinstance(v, Decimal) else v for v in r.get("values", r["cells"])])
            if r.get("kind") in BOLD_KINDS:
                for c in ws[ws.max_row]:
                    c.font = Font(bold=True)
            if r.get("indent"):
                ws.cell(ws.max_row, 1).alignment = Alignment(indent=1)
        for ci in sh["num_cols"]:
            col = get_column_letter(ci + 1)
            for cell in ws[col][5:]:
                cell.number_format = NUM_FMT
                cell.alignment = Alignment(horizontal="right")
        for i, w in enumerate(sh["widths"]):
            ws.column_dimensions[get_column_letter(i + 1)].width = w
        if k == 0 and notes:
            ws.append([])
            for n in notes:
                ws.append([n])
        ws.freeze_panes = "A6"
    wb.save(out_path)
    return str(out_path)


def _esc(s) -> str:
    return (str(s or "")).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def _n(d) -> str:
    return fmt(d) if d else "-"


# ---------------------------------------------------------------------------------------- #
# Trial Balance
# ---------------------------------------------------------------------------------------- #
def trial_balance(ledger: Ledger, fy: int, out_path: str | Path, upto: dt.date | None = None,
                  as_excel: bool = False) -> str:
    tb = ledger.trial_balance(fy, upto)
    header = ["Code", "Account", "Account type", "Report code", "Debit", "Credit"]
    rows = [{"cells": [r["code"], r["name"], r["type"], r["report_code"], _n(r["debit"]), _n(r["credit"])],
             "values": [r["code"], r["name"], r["type"], r["report_code"], r["debit"], r["credit"]]}
            for r in tb["rows"]]
    rows.append({"kind": "total", "cells": ["", "Total", "", "", _n(tb["total_debit"]), _n(tb["total_credit"])],
                 "values": ["", "Total", "", "", tb["total_debit"], tb["total_credit"]]})
    period = f"As at {tb['as_at'].day} {tb['as_at']:%B %Y}"
    ledger.store.audit(ledger.slug, "report.generated", {"fy": fy, "report": "trial_balance", "file": Path(out_path).name})
    if as_excel:
        return _xlsx(out_path, "Trial Balance", ledger.client.name, period, header, rows, {4, 5}, [12, 44, 20, 22, 16, 16])
    return _pdf_table(out_path, "Trial Balance", ledger.client.name, period, header, rows,
                      [16 * mm, 68 * mm, 28 * mm, 30 * mm, 17 * mm, 17 * mm], {4, 5})


# ---------------------------------------------------------------------------------------- #
# Profit and Loss, Balance Sheet
# ---------------------------------------------------------------------------------------- #
def _statement_rows(data: dict) -> list[dict]:
    rows = []
    for r in data["rows"]:
        vals = [r["cy"]] + ([r["py"]] if data["compare"] else [])
        heading = r["kind"] in ("heading", "subheading")
        cells = [r["label"]] + ["" if heading or v is None else _n(v) for v in vals]
        rows.append({"kind": r["kind"], "cells": cells, "values": [r["label"]] + [None if heading else v for v in vals],
                     "indent": 3 if r["kind"] == "line" else 0})
    return rows


def _statement_file(ledger: Ledger, data: dict, key: str, out_path, as_excel: bool) -> str:
    header = [""] + data["columns"]
    rows = _statement_rows(data)
    n = len(data["columns"])
    ledger.store.audit(ledger.slug, "report.generated", {"report": key, "file": Path(out_path).name,
                                                         "period": data["period"]})
    if as_excel:
        return _xlsx(out_path, data["title"], ledger.client.name, data["period"], header, rows,
                     set(range(1, n + 1)), [52] + [18] * n)
    return _pdf_table(out_path, data["title"], ledger.client.name, data["period"], header, rows,
                      [(176 - 30 * n) * mm] + [30 * mm] * n, set(range(1, n + 1)))


def profit_and_loss(ledger: Ledger, out_path: str | Path, start: dt.date, end: dt.date, compare: bool = True,
                    as_excel: bool = False) -> str:
    return _statement_file(ledger, stm.profit_and_loss(ledger, start, end, compare), "profit_and_loss",
                           out_path, as_excel)


def balance_sheet(ledger: Ledger, out_path: str | Path, as_at: dt.date, compare: bool = True,
                  as_excel: bool = False) -> str:
    return _statement_file(ledger, stm.balance_sheet(ledger, as_at, compare), "balance_sheet", out_path, as_excel)


# ---------------------------------------------------------------------------------------- #
# GST report
# ---------------------------------------------------------------------------------------- #
def gst_notes(data: dict) -> list[str]:
    notes = ["Cash basis: GST on money received and paid in the period - the bank transactions." if data["basis"] == "cash"
             else "Accrual basis: GST on bank transactions and on journals with a tax rate, by their date."]
    if data["left_out"]:
        notes.append(f"{data['left_out']} journal line(s) with a tax rate are left out: on the cash basis only money "
                     f"received and paid counts.")
    if not data["registered"]:
        notes.append("The client is not marked as registered for GST (Client details).")
    return notes


def gst_report(ledger: Ledger, out_path: str | Path, start: dt.date, end: dt.date, basis: str = "accrual",
               as_excel: bool = False) -> str:
    data = stm.gst_report(ledger, start, end, basis)
    period = f"{data['period']} - {data['basis_label']}"
    bas = [{"cells": [k, label, fmt(v, blank_zero=False)], "values": [k, label, v]} for k, label, v in data["labels"]]
    net = data["net_gst"]
    bas.append({"kind": "grand", "cells": ["", "GST payable" if net >= 0 else "GST refundable", fmt(abs(net), blank_zero=False)],
                "values": ["", "GST payable" if net >= 0 else "GST refundable", abs(net)]})
    rates = [{"cells": [r["tax"], r["side"].capitalize(), fmt(r["net"], blank_zero=False), fmt(r["gst"], blank_zero=False),
                        fmt(r["gross"], blank_zero=False)],
              "values": [r["tax"], r["side"].capitalize(), r["net"], r["gst"], r["gross"]]} for r in data["rates"]]
    detail = []
    for tax in dict.fromkeys(x["tax"] for x in data["lines"]):
        xs = [x for x in data["lines"] if x["tax"] == tax]
        detail.append({"kind": "subheading", "cells": [tax, "", "", "", "", "", ""], "values": [tax, None, None, None, None, None, None]})
        for x in xs:
            acct = f"{x['account']} {x['account_name']}"
            detail.append({"cells": [f"{x['date']:%d/%m/%Y}", x["number"], acct, x["text"], _n(x["net"]), _n(x["gst"]),
                                     _n(x["gross"])],
                           "values": [x["date"], x["number"], acct, x["text"], x["net"], x["gst"], x["gross"]]})
        t = [sum((x[k] for x in xs), ZERO) for k in ("net", "gst", "gross")]
        detail.append({"kind": "total", "cells": ["", "", "", f"Total {tax}"] + [_n(v) for v in t],
                       "values": ["", "", "", f"Total {tax}"] + t})
    if not rates:
        rates = [{"cells": ["No transactions with GST in this period", "", "", "", ""], "values": ["No transactions with GST in this period"]}]
    notes = gst_notes(data)
    ledger.store.audit(ledger.slug, "report.generated", {"report": "gst_report", "basis": basis,
                                                         "file": Path(out_path).name, "period": data["period"]})
    if as_excel:
        return _xlsx_sheets(out_path, "GST Report", ledger.client.name, period, [
            {"title": "BAS summary", "header": ["Label", "Description", "Amount"], "rows": bas, "num_cols": {2},
             "widths": [10, 50, 18]},
            {"title": "By tax rate", "header": ["Tax rate", "Sales or purchases", "Net", "GST", "Gross"], "rows": rates,
             "num_cols": {2, 3, 4}, "widths": [24, 20, 16, 16, 16]},
            {"title": "Transactions", "header": ["Date", "Number", "Account", "Description", "Net", "GST", "Gross"],
             "rows": detail, "num_cols": {4, 5, 6}, "widths": [12, 13, 34, 50, 16, 16, 16]}], notes)
    return _pdf_tables(out_path, "GST Report", ledger.client.name, period, [
        {"heading": "BAS summary", "header": ["Label", "Description", "Amount"], "rows": bas,
         "widths": [20 * mm, 126 * mm, 30 * mm], "num_cols": {2}},
        {"heading": "By tax rate", "header": ["Tax rate", "Sales or purchases", "Net", "GST", "Gross"], "rows": rates,
         "widths": [50 * mm, 36 * mm, 30 * mm, 30 * mm, 30 * mm], "num_cols": {2, 3, 4}},
        {"heading": "Transactions", "header": ["Date", "Number", "Account", "Description", "Net", "GST", "Gross"],
         "rows": detail, "widths": [17 * mm, 20 * mm, 38 * mm, 47 * mm, 18 * mm, 18 * mm, 18 * mm],
         "num_cols": {4, 5, 6}}], notes=notes)


# ---------------------------------------------------------------------------------------- #
# General Ledger Detail (account transactions)
# ---------------------------------------------------------------------------------------- #
def general_ledger(ledger: Ledger, fy: int, out_path: str | Path, start: dt.date | None = None,
                   end: dt.date | None = None, codes: list[str] | None = None, as_excel: bool = False) -> str:
    data = ledger.general_ledger(fy, start, end, codes)
    s, e = ledger.fy_range(fy)
    start, end = max(start or s, s), min(end or e, e)
    header = ["Date", "Journal", "Narration / description", "Debit", "Credit", "Balance"]
    rows = []
    for acc in data:
        rows.append({"kind": "heading", "cells": ["", "", f"{acc['code']} {acc['name']}", "", "", ""],
                     "values": ["", "", f"{acc['code']} {acc['name']}", None, None, None]})
        rows.append({"cells": ["", "", "Opening balance", "", "", _n(acc["opening"])],
                     "values": ["", "", "Opening balance", None, None, acc["opening"]]})
        for x in acc["entries"]:
            text = x["narration"] + (f" - {x['description']}" if x["description"] else "")
            rows.append({"cells": [f"{x['date']:%d/%m/%Y}", x["number"], text, _n(x["debit"]), _n(x["credit"]),
                                   _n(x["balance"])],
                         "values": [x["date"], x["number"], text, x["debit"], x["credit"], x["balance"]]})
        rows.append({"kind": "total", "cells": ["", "", f"Total {acc['code']} {acc['name']}", _n(acc["debits"]),
                                                _n(acc["credits"]), _n(acc["closing"])],
                     "values": ["", "", f"Total {acc['code']} {acc['name']}", acc["debits"], acc["credits"], acc["closing"]]})
    period = f"{start.day} {start:%B %Y} to {end.day} {end:%B %Y}"
    ledger.store.audit(ledger.slug, "report.generated", {"fy": fy, "report": "general_ledger", "file": Path(out_path).name})
    if as_excel:
        return _xlsx(out_path, "General Ledger Detail", ledger.client.name, period, header, rows, {3, 4, 5},
                     [12, 13, 60, 16, 16, 16])
    return _pdf_table(out_path, "General Ledger Detail", ledger.client.name, period, header, rows,
                      [18 * mm, 20 * mm, 76 * mm, 21 * mm, 21 * mm, 21 * mm], {3, 4, 5})


# ---------------------------------------------------------------------------------------- #
# Journal Report
# ---------------------------------------------------------------------------------------- #
def journal_report(ledger: Ledger, fy: int, out_path: str | Path, include_void: bool = True,
                   as_excel: bool = False) -> str:
    statuses = ("posted", "void") if include_void else ("posted",)
    header = ["Date", "Journal", "Account", "Description", "Debit", "Credit"]
    rows = []
    for j in ledger.journals(fy, statuses=statuses):
        status = " (VOID)" if j.status == "void" else ""
        rows.append({"kind": "heading", "cells": [f"{j.date:%d/%m/%Y}", j.number, j.narration + status, j.kind, "", ""],
                     "values": [j.date, j.number, j.narration + status, j.kind, None, None]})
        for ln in j.lines:
            a = ledger.account(ln.account)
            name = f"{ln.account} {a.name}" if a else ln.account
            rows.append({"cells": ["", "", name, ln.description, _n(ln.debit), _n(ln.credit)],
                         "values": ["", "", name, ln.description, ln.debit, ln.credit]})
        rows.append({"kind": "total", "cells": ["", "", "", "Total", _n(j.total_debit), _n(j.total_credit)],
                     "values": ["", "", "", "Total", j.total_debit, j.total_credit]})
    s, e = ledger.fy_range(fy)
    period = f"{s.day} {s:%B %Y} to {e.day} {e:%B %Y}"
    ledger.store.audit(ledger.slug, "report.generated", {"fy": fy, "report": "journal_report", "file": Path(out_path).name})
    if as_excel:
        return _xlsx(out_path, "Journal Report", ledger.client.name, period, header, rows, {4, 5}, [12, 13, 44, 40, 16, 16])
    return _pdf_table(out_path, "Journal Report", ledger.client.name, period, header, rows,
                      [18 * mm, 20 * mm, 58 * mm, 46 * mm, 17 * mm, 17 * mm], {4, 5})


REPORTS = {
    "financial_reports": "Financial Reports pack",
    "trial_balance": "Trial Balance",
    "profit_and_loss": "Profit and Loss",
    "balance_sheet": "Balance Sheet",
    "general_ledger": "General Ledger Detail",
    "gst_report": "GST Report",
    "journal_report": "Journal Report",
}
