#!/usr/bin/env python3
"""
build_pack.py - post the workpaper's Profit & Loss and Balance Sheet tabs to the
firm's Xero chart of accounts and write the ledger pack (pack.json) that
render_pack.py prints.

This is the "Xero ledger" step.  The entity type picks the chart
(assets/chart-of-accounts/company.csv | trust.csv | partnership.csv |
sole-trader.csv); every workpaper line is then posted to the chart account that
matches its nature (scripts/chart.py: section header + keyword rules + name
similarity).  The account's Xero report code decides where it prints.

Usage
-----
  python3 scripts/build_pack.py --workpaper "2026 Workpapers - Client Pty Ltd.xlsm" \\
      --name "Client Pty Ltd" --abn "12 345 678 901" --fy-end 2026-06-30 --type company \\
      --directors "Jane Citizen;John Citizen" --out pack.json --mapping-out mapping.md

  # change where a line prints (repeatable; key = workpaper code or account name)
      --set-code "650.1=ASS.NCA.REC" --set-code "Paid Parental Leave=LIA.CUR.PAY.EMP"
  # post a line to a specific chart account (its report code applies)
      --set-account "Loan - Director=31603"
  # print a line under a different label
      --rename "Cash on Hands=Cash on Hand"
  # re-apply every override saved in last run's pack (or a standalone overrides JSON)
      --overrides pack.json

  # other options
      --pl / --bs CSV (label,cy,py) instead of --workpaper
      --chart PATH                    explicit chart CSV (default: bundled chart for --type)
      --depreciation depn.csv|json    (references/report-pack-structure.md)
      --opening-equity-py 85089       PY opening equity from last year's pack
      --expenses-negative auto|yes|no
      --ledger-out ledger.csv         the workpaper as a trial balance on the chart

The script prints the mapping table and the validation report.  Lines marked
REVIEW must be resolved (override or chart change) before the pack is rendered.
"""
from __future__ import annotations

import argparse
import csv
import datetime as _dt
import json
import re
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from chart import SECTIONS, Chart, Matcher, canonical_type, detect_entity_type, normalise, tokens  # noqa: E402
from fr_engine import Pack, describe_code, fmt, load_map, resolve_slot, strip_code  # noqa: E402

SKIP_LABELS = re.compile(
    r"^(total\b|gross profit|net profit|net loss|operating profit|net assets|net income|profit/\(loss\)|"
    r"profit before|profit after|less |plus |add |ebit|net (operating )?(profit|loss|income)|"
    r"(undistributed|retained) (profit|income|earnings) (for|available)|profit available|balance sheet difference)", re.I)
APPORTION = re.compile(r"\s*\(\d+(\.\d+)?\s*%[^)]*\)\s*$")      # "Cleaning (11.3% Office Area)" -> "Cleaning"
PERIOD_HDR = re.compile(r"^((19|20)\d{2}|\d{1,2}\s+[A-Za-z]+\s+(19|20)\d{2})$")
HEADER_HINTS = re.compile(r"^(account( code)?( - name)?|particulars|cy|py|current year|prior year|description|\d{4})$", re.I)
EXTRA_HEADERS = {"income", "expenses", "assets", "liabilities", "equity"}


# --------------------------------------------------------------------------- #
# reading the two statements
# --------------------------------------------------------------------------- #
def _num(v):
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace(",", "").replace("$", "")
    neg = s.startswith("(") and s.endswith(")")
    s = s.strip("()")
    if s in ("-", "", "–"):
        return 0.0
    try:
        f = float(s)
    except ValueError:
        return None
    return -f if neg else f


def read_rows_xlsx(path: str, sheet_names: list[str]) -> list[list]:
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    for name in sheet_names:
        if name in wb.sheetnames:
            ws = wb[name]
            return [list(r) for r in ws.iter_rows(values_only=True)]
    raise SystemExit(f"None of the sheets {sheet_names} found in {path}. Sheets: {wb.sheetnames}")


def _is_period_header(v) -> bool:
    if isinstance(v, (_dt.date, _dt.datetime)):
        return True
    if isinstance(v, (int, float)) and float(v).is_integer() and 1900 < v < 2100:
        return True
    return isinstance(v, str) and bool(PERIOD_HDR.match(v.strip()))


def split_side_by_side(rows: list[list]) -> tuple[list[list], list[list]]:
    """
    An ITR workpaper 'Business' tab holds the pasted Xero P&L (left) and Balance Sheet (right) side by side,
    e.g. P&L in A-C and BS in F-H (2025) or K-N (2026).  The BS block starts at the 'Balance Sheet' title;
    in each block the first year / date header is CY, the column right after it PY (when it is one too),
    and the label is the first text cell left of the CY column.  Returns [label, cy, py] rows per block.
    """
    width = max((len(r) for r in rows), default=0)
    b0 = next((j for r in rows[:10] for j, v in enumerate(r)
               if j > 0 and isinstance(v, str) and v.strip().lower() == "balance sheet"), None)
    if b0 is None:
        raise SystemExit("No 'Balance Sheet' title to the right of the P&L on this tab - use --pl-sheet / --bs-sheet.")

    def block(lo: int, hi: int) -> list[list]:
        cy = py = None
        for r in rows[:15]:
            cols = [j for j in range(lo, min(hi, len(r))) if _is_period_header(r[j])]
            if cols:
                cy = cols[0]
                py = cols[1] if len(cols) > 1 and cols[1] == cy + 1 else None
                break
        if cy is None:
            raise SystemExit(f"No year / date column header found in columns {lo + 1}-{hi} of the Business tab.")
        out = []
        for r in rows:
            r = list(r) + [None] * (width - len(r))
            label = next((str(r[j]).strip() for j in range(lo, cy) if isinstance(r[j], str) and r[j].strip()), "")
            out.append([label, r[cy], r[py] if py is not None else None])
        return out

    return block(0, b0), block(b0, width)


def read_rows_csv(path: str) -> list[list]:
    with open(path, newline="", encoding="utf-8-sig") as fh:
        return [row for row in csv.reader(fh)]


def parse_statement(rows: list[list], side: str, stop_at: str | None = None) -> list[dict]:
    """
    Walk the rows of a P&L or BS tab and return line dicts {code, name, section, cy, py, row}.
    A label with no figures is a section header if it is one of chart.SECTIONS; totals and
    computed lines are skipped - the engine recomputes them.
    """
    out = []
    section = None
    started = False
    net = [0.0, 0.0]          # running P&L result of the lines read so far (income less expenses)
    for idx, r in enumerate(rows, 1):
        r = list(r) + [None] * (4 - len(r))
        label = (str(r[0]).strip() if r[0] is not None else "")
        cy, py = _num(r[1]), _num(r[2])
        if not label:
            continue
        key = label.lower().rstrip(":").strip()
        if not started:
            if HEADER_HINTS.match(label) or (isinstance(r[1], str) and HEADER_HINTS.match(str(r[1]).strip())):
                started = True
                continue
            if cy is None and py is None:
                if key in SECTIONS:
                    started = True
                    section = key
                continue
            started = True
        if stop_at and re.match(stop_at, label, re.I):
            break             # e.g. ITR workpapers: tax adjustments follow 'Net Profit'
        if cy is None and py is None:
            if key in SECTIONS or key in EXTRA_HEADERS:
                section = key
            continue
        if SKIP_LABELS.match(label) and not re.match(r"^\d", label):
            continue
        code, name = strip_code(label)
        name = APPORTION.sub("", name)
        if side == "pl" and not code and re.search(r"profit|loss|surplus|deficit|result", key) \
                and abs((cy or 0.0) - net[0]) < 0.5 and abs((py or 0.0) - net[1]) < 0.5:
            continue          # an un-coded subtotal (e.g. a mistyped 'Opearting Profit') - the engine recomputes it
        if side == "pl":
            k = {"cr": 1, "dr": -1}.get((SECTIONS.get(section) or (None, None, None))[2], 0)
            net[0] += k * (cy or 0.0)
            net[1] += k * (py or 0.0)
        out.append({"code": code, "name": name.strip(), "section": section, "cy": cy or 0.0, "py": py or 0.0,
                    "row": idx, "side": side})
    return out


# --------------------------------------------------------------------------- #
# overrides
# --------------------------------------------------------------------------- #
def _kv(items: list[str] | None, what: str) -> dict:
    out = {}
    for it in items or []:
        if "=" not in it:
            raise SystemExit(f"{what} expects KEY=VALUE, got '{it}'")
        k, v = it.split("=", 1)
        out[k.strip()] = v.strip()
    return out


def load_overrides(path: str | None) -> dict:
    if not path:
        return {}
    data = json.loads(Path(path).read_text(encoding="utf-8"))
    return data.get("overrides", data) if isinstance(data, dict) else {}


def find_override(table: dict, line: dict) -> str | None:
    """Override keys match the workpaper code, the account name, or 'code - name' (case-insensitive)."""
    if not table:
        return None
    keys = {k.strip().lower(): v for k, v in table.items()}
    cands = [line.get("code") or "", line["name"], f"{line.get('code')} - {line['name']}" if line.get("code") else ""]
    for c in cands:
        if c and c.strip().lower() in keys:
            return keys[c.strip().lower()]
    n = normalise(line["name"])
    for k, v in keys.items():
        if normalise(k) == n:
            return v
    return None


# --------------------------------------------------------------------------- #
# posting to the chart
# --------------------------------------------------------------------------- #
def nature_of_slot(slot: str) -> str:
    """Natural sign of a layout slot: 'dr' amounts are positive debits, 'cr' positive credits."""
    if slot.startswith(("EXP", "ASS")) or slot in ("EQU.DIV",):
        return "dr"
    return "cr"


def post_lines(lines: list[dict], chart: Chart, rmap: dict, overrides: dict, expenses_negative: str,
               entity_type: str | None = None) -> list[dict]:
    matcher = Matcher(chart, entity_type)
    ov_codes = overrides.get("report_codes") or {}
    ov_accts = overrides.get("accounts") or {}
    ov_labels = overrides.get("labels") or {}
    accounts = []
    for ln in lines:
        side = ln["side"]
        sec_side, section, nature = SECTIONS.get((ln["section"] or "").lower(), (None, None, None))
        m = matcher.match(ln["name"], ln["section"], ln.get("code"), side_hint=side)
        acct, rcode, conf, why = m.account, m.report_code, m.confidence, m.reason
        forced_acct = find_override(ov_accts, ln)
        if forced_acct:
            a2 = chart.by_code(forced_acct)
            if not a2:
                raise SystemExit(f"--set-account: chart {chart.path.name} has no account {forced_acct}")
            acct, rcode, conf, why = a2, a2.report_code, "override", f"posted to chart account {a2.short()} by override"
        forced_code = find_override(ov_codes, ln)
        if forced_code:
            rcode = forced_code.upper()
            if not forced_acct:
                # keep the ledger consistent: post to the chart account that carries the new code
                same = [a for a in chart.in_family(["=" + rcode]) if not a.system]
                acct = matcher._best(tokens(ln["name"]), same)[0] if same else None
            conf, why = "override", f"report code set to {rcode} by override" + (
                f" -> chart account {acct.short()}" if acct else " (layout code - no chart account carries it)")
        review = m.review and conf != "override"
        if entity_type == "sole_trader" and side == "bs" and nature == "cr" and not (forced_acct or forced_code) \
                and (section or "").endswith("liabilities") and (ln["cy"] < 0 or (ln["cy"] == 0 and ln["py"] < 0)) \
                and re.search(r"owner|proprietor|funds? (introduced|contributed)|drawings|\bgst\b", ln["name"], re.I):
            # the firm's sole-trader Xero pack prints these debit balances under Receivables > Current
            if re.search(r"\bgst\b", ln["name"], re.I):
                rcode = "ASS.CUR.REC"
            else:
                cands = [a for a in chart.in_family(["=ASS.CUR.LOA"]) if not a.system]
                acct = matcher._best(tokens(ln["name"]), cands)[0] if cands else None
                rcode = acct.report_code if acct else "ASS.CUR.LOA"
            conf, review = "rule", False
            why = (f"debit balance on a liability account - printed under Receivables > Current, as the firm's "
                   f"sole-trader Xero pack does ({rcode}" + (f", ledger {acct.short()})" if acct else ")"))
        slot, person, _ = resolve_slot(rcode, ln["name"], rmap)
        label = find_override(ov_labels, ln)
        rec = {
            "code": ln.get("code"), "name": ln["name"],
            "type": acct.type if acct else "", "report_code": rcode,
            "chart_code": acct.code if acct else "", "chart_name": acct.name if acct else "",
            "cy": ln["cy"], "py": ln["py"], "source": f"{side}:row{ln['row']}",
            "mapping": {"confidence": conf, "reason": why, "section": ln["section"],
                        "prints_as": describe_code(rcode, ln["name"], rmap),
                        "review": review},
        }
        if label:
            rec["display_name"] = label
        if acct and acct.reporting_name and not label:
            rec["chart_reporting_name"] = acct.reporting_name
        if person is not None:
            rec["person"] = person
        rec["_nature_src"] = nature
        rec["_slot"] = slot or ""
        accounts.append(rec)

    # expenses stored as negatives in the workpaper (auto: most expense-section lines negative)
    exp = [a for a in accounts if a["_nature_src"] == "dr" and a["source"].startswith("pl")]
    if exp:
        neg = sum(1 for a in exp if a["cy"] < 0 or (a["cy"] == 0 and a["py"] < 0))
        flip = expenses_negative == "yes" or (expenses_negative == "auto" and neg > len(exp) / 2)
        if flip:
            for a in exp:
                a["cy"], a["py"] = -a["cy"], -a["py"]
                a["mapping"]["reason"] += "; sign flipped (workpaper shows expenses as negatives)"

    # a line printed on the other side from where the workpaper shows it changes sign
    # (e.g. a loss on disposal under Expenses coded to Other Income, dividends in equity shown as paid)
    for a in accounts:
        src = a.pop("_nature_src")
        slot = a.pop("_slot")
        if not src or not slot:
            continue
        tgt = nature_of_slot(slot)
        if src != tgt:
            a["cy"], a["py"] = -a["cy"], -a["py"]
            a["mapping"]["reason"] += f"; sign reversed ({'debit' if src == 'dr' else 'credit'} line printed in a {'debit' if tgt == 'dr' else 'credit'} position)"
    return accounts


def mapping_markdown(accounts: list[dict]) -> str:
    out = ["| Source | WP code | Workpaper account | CY | Chart account | Report code | Prints as | Confidence | Why |",
           "|---|---|---|---:|---|---|---|---|---|"]
    for a in accounts:
        m = a["mapping"]
        flag = " **REVIEW**" if m.get("review") else ""
        chart = f"{a['chart_code']} {a['chart_name']}".strip() or "-"
        out.append(f"| {a['source']} | {a['code'] or ''} | {a['name']} | {fmt(a['cy'])} | {chart} | `{a['report_code']}`{flag} | "
                   f"{m.get('prints_as', '')} | {m['confidence']} | {m['reason']} |")
    return "\n".join(out)


def write_ledger(accounts: list[dict], path: str):
    with open(path, "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh)
        w.writerow(["Chart Code", "Chart Account", "Report Code", "Account Type", "Workpaper Code", "Workpaper Account",
                    "CY", "PY", "Prints As", "Confidence"])
        for a in accounts:
            w.writerow([a["chart_code"], a["chart_name"], a["report_code"], a["type"], a["code"] or "", a["name"],
                        round(a["cy"], 2), round(a["py"], 2), a["mapping"].get("prints_as", ""), a["mapping"]["confidence"]])


# --------------------------------------------------------------------------- #
# depreciation input
# --------------------------------------------------------------------------- #
def read_depreciation(path: str) -> dict:
    p = Path(path)
    if p.suffix.lower() == ".json":
        return json.loads(p.read_text(encoding="utf-8"))
    groups: dict[str, list] = {}
    with open(p, newline="", encoding="utf-8-sig") as fh:
        for row in csv.DictReader(fh):
            row = {k.strip().lower().replace(" ", "_"): (v or "").strip() for k, v in row.items() if k}
            g = row.pop("group", None) or row.pop("class", None) or "Assets"
            asset = {}
            for k, v in row.items():
                if k in ("cost", "opening_accum", "opening_value", "purchases", "disposals", "depreciation",
                         "closing_accum", "closing_value"):
                    asset[k] = _num(v) or 0.0
                elif v:
                    asset[k] = v
            groups.setdefault(g, []).append(asset)
    layout = "simple"
    cols = set().union(*[set(a) for g in groups.values() for a in g]) if groups else set()
    if "asset_number" in cols or "opening_accum" in cols:
        layout = "register"
    elif "rate" in cols or "method" in cols:
        layout = "standard"
    return {"layout": layout, "groups": [{"name": g, "assets": a} for g, a in groups.items()]}


# --------------------------------------------------------------------------- #
def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    src = ap.add_argument_group("source")
    src.add_argument("--workpaper", help="xlsx/xlsm workpaper")
    src.add_argument("--pl-sheet", default="Profit & Loss,PL")
    src.add_argument("--bs-sheet", default="Balance Sheet,BS")
    src.add_argument("--business-tab", nargs="?", const="Business", default=None,
                     help="read the P&L and Balance Sheet pasted side by side on one tab (ITR workpapers: "
                          "'Business'); used automatically when the workpaper has no P&L tab but has 'Business'")
    src.add_argument("--pl", help="CSV: label,cy,py")
    src.add_argument("--bs", help="CSV: label,cy,py")
    ent = ap.add_argument_group("entity")
    ent.add_argument("--entity", help="JSON file with entity fields (name, abn, fy_end, type, directors, trustee, partners ...)")
    ent.add_argument("--name"); ent.add_argument("--abn"); ent.add_argument("--fy-end")
    ent.add_argument("--type", help="company | trust | partnership | sole_trader (default: detected)")
    ent.add_argument("--directors", "--signatories", dest="signatories",
                     help="semicolon separated: directors (company / corporate trustee), trustees, partners or the proprietor")
    ent.add_argument("--trustee", help="trustee name (trusts), e.g. 'ABC Pty Ltd' or 'John Citizen'")
    ent.add_argument("--business-name", help="trading name (sole traders / partnerships)")
    ent.add_argument("--trustee-type", choices=["corporate", "individual"],
                     help="trusts: corporate trustee (directors sign) or individual trustee(s) (default: from the trustee name)")
    ent.add_argument("--legal-name", help="name used in the Compilation Report when it differs from the cover name")
    ent.add_argument("--proprietor", help="sole traders: the proprietor's name (Note 1 / declaration wording)")
    ent.add_argument("--date-signed", help="date printed beside the signatures, e.g. '5 August 2026' (blank by default)")
    ent.add_argument("--compiled-by", help="compilation report signatory (default: boilerplate firm signatory)")
    mp = ap.add_argument_group("mapping")
    mp.add_argument("--chart", "--coa", dest="chart", help="chart CSV (default: assets/chart-of-accounts/<type>.csv)")
    mp.add_argument("--set-code", action="append", metavar="KEY=REPORT_CODE")
    mp.add_argument("--set-account", action="append", metavar="KEY=CHART_CODE")
    mp.add_argument("--rename", action="append", metavar="KEY=LABEL")
    mp.add_argument("--overrides", help="overrides JSON, or a previous pack.json (its 'overrides' block)")
    ap.add_argument("--map")
    ap.add_argument("--depreciation")
    ap.add_argument("--opening-equity-py", type=float)
    ap.add_argument("--expenses-negative", choices=["auto", "yes", "no"], default="auto")
    ap.add_argument("--cost-of-sales-label", help="heading for the cost of sales section (default by entity type: "
                    "'Cost of Goods Sold' for companies and partnerships, 'Cost of Sales' otherwise)")
    ap.add_argument("--cos-lines", choices=["accounts", "grouped"],
                    help="cost of sales lines: one per account, or grouped as Opening inventory / Purchases / Closing "
                         "inventory (default: grouped for partnerships, accounts otherwise)")
    ap.add_argument("--sort", choices=["alpha", "source"], default="alpha")
    ap.add_argument("--out", required=True)
    ap.add_argument("--mapping-out")
    ap.add_argument("--ledger-out")
    args = ap.parse_args()

    stop_at = None
    if args.workpaper:
        import openpyxl
        names = openpyxl.load_workbook(args.workpaper, read_only=True).sheetnames
        pl_names = [x.strip() for x in args.pl_sheet.split(",")]
        if args.business_tab is None and not any(n in names for n in pl_names) and "Business" in names:
            args.business_tab = "Business"     # ITR workpaper (sole traders): no P&L / BS tabs
        if args.business_tab:
            pl_rows, bs_rows = split_side_by_side(read_rows_xlsx(args.workpaper, [args.business_tab]))
            stop_at = r"^net (profit|loss)\b"
        else:
            pl_rows = read_rows_xlsx(args.workpaper, pl_names)
            bs_rows = read_rows_xlsx(args.workpaper, [x.strip() for x in args.bs_sheet.split(",")])
    elif args.pl and args.bs:
        pl_rows, bs_rows = read_rows_csv(args.pl), read_rows_csv(args.bs)
    else:
        ap.error("give --workpaper or both --pl and --bs")

    entity = {}
    if args.entity:
        entity = json.loads(Path(args.entity).read_text(encoding="utf-8"))
    for k, v in (("name", args.name), ("abn", args.abn), ("fy_end", args.fy_end), ("trustee", args.trustee),
                 ("business_name", args.business_name), ("trustee_type", args.trustee_type),
                 ("legal_name", args.legal_name), ("proprietor", args.proprietor), ("date_signed", args.date_signed),
                 ("compiled_by", args.compiled_by)):
        if v:
            entity[k] = v
    if args.signatories:
        entity["directors"] = [d.strip() for d in args.signatories.split(";") if d.strip()]

    pl = parse_statement(pl_rows, "pl", stop_at)
    bs = parse_statement(bs_rows, "bs")

    etype = args.type or entity.get("type")
    if etype:
        etype = canonical_type(etype)
        why_type = "given"
    else:
        etype, why_type = detect_entity_type(entity.get("name", ""), [l["name"] for l in bs if (l["section"] or "") == "equity"])
        if etype == "smsf":
            raise SystemExit("This looks like a superannuation fund - SMSFs are out of scope (use the SMSF audit pack).")
    entity["type"] = etype

    rmap = load_map(args.map)
    chart = Chart(args.chart, etype) if args.chart else Chart.for_type(etype)
    overrides = load_overrides(args.overrides)
    for key, items in (("report_codes", args.set_code), ("accounts", args.set_account), ("labels", args.rename)):
        new = _kv(items, f"--{key}")
        if new:
            overrides.setdefault(key, {}).update(new)

    accounts = post_lines(pl + bs, chart, rmap, overrides, args.expenses_negative, etype)

    pack = {
        "_generated_by": "build_pack.py",
        "_source": {"workpaper": args.workpaper, "pl": args.pl, "bs": args.bs, "chart": str(chart.path.name),
                    "entity_type": f"{etype} ({why_type})"},
        "entity": entity,
        "options": {"sort": args.sort,
                    **({"cost_of_sales_label": args.cost_of_sales_label} if args.cost_of_sales_label else {}),
                    **({"cos_lines": args.cos_lines} if args.cos_lines else {})},
        "overrides": overrides,
        "accounts": accounts,
    }
    if args.opening_equity_py is not None:
        pack["equity"] = {"opening_py": args.opening_equity_py}
    if args.depreciation:
        pack["depreciation"] = read_depreciation(args.depreciation)

    Path(args.out).write_text(json.dumps(pack, indent=2, ensure_ascii=False), encoding="utf-8")
    md = mapping_markdown(accounts)
    if args.mapping_out:
        Path(args.mapping_out).write_text(f"# Account mapping - {entity.get('name', '')}\n\nChart: {chart.path.name} "
                                          f"({etype}, {why_type})\n\n{md}\n", encoding="utf-8")
    if args.ledger_out:
        write_ledger(accounts, args.ledger_out)

    print(f"Entity type: {etype} ({why_type}); chart: {chart.path.name}\n")
    print("## Account mapping\n")
    print(md)
    review = [a for a in accounts if a["mapping"].get("review")]
    print(f"\n{len(accounts)} lines posted; {len(review)} need review.")
    for a in review:
        print(f"  REVIEW  {a['name']}  ->  {a['report_code']}  ({a['mapping']['reason']})")

    print("\n## Validation\n")
    p = Pack(pack, rmap)
    print(p.validation_markdown())
    s = p.summary()
    print(f"\nNPAT {fmt(s['npat']['cy'])} (PY {fmt(s['npat']['py'])}) | Net Assets {fmt(s['net_assets']['cy'])} "
          f"(PY {fmt(s['net_assets']['py'])}) | Notes: {', '.join(f'{n}. {t}' for n, t in s['notes'])}")
    print(f"\nWrote {args.out}")


if __name__ == "__main__":
    main()
