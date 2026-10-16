"""
Imports - CSV / Excel files into the ledger.

* Balances ("classified account code data"): a trial balance or account list with codes, names and
  amounts (debit / credit columns, one signed balance column, or CY and PY columns), optionally with
  report codes or Xero account types.  Xero trial-balance exports are recognised.  A balances import
  *sets* the accounts to the imported figures: it posts the difference from what the ledger already
  holds (e.g. balances carried forward from last year), so loading each year's trial balance never
  double counts.
* Journals: Xero's manual-journal import CSV (*Narration, *Date, *AccountCode, *Amount ...) or a plain
  Date / Reference / Narration / Account / Debit / Credit file.
* Chart of accounts: Xero's chart export (handled by ClientChart.merge).
* Workpaper: the Fortis year-end workpaper (Profit & Loss / Balance Sheet tabs, or an ITR Business tab),
  read and classified by the financial-report-builder skill.

Every import is planned first (what each line will post to, new accounts, rounding) so the user can
check it, then applied.
"""
from __future__ import annotations

import csv
import io
import re
from dataclasses import asdict, dataclass, field
from decimal import Decimal
from pathlib import Path

from .chartbook import ChartError, code_sort_key, default_report_code, is_pl_account
from .engine import REPORT_MAP, build_pack, chartlib, describe, fr_engine, nature, slot_of
from .ledger import Ledger, LedgerError, ValidationError
from .model import Journal, Line, parse_date
from .money import ZERO, as_str, is_amount, to_dec

ROUNDING_TOLERANCE = Decimal("1.00")

XERO_TYPE_SECTION = {
    "bank": "Bank", "current asset": "Current Assets", "inventory": "Current Assets", "prepayment": "Current Assets",
    "fixed asset": "Fixed Assets", "non-current asset": "Non-current Assets", "current liability": "Current Liabilities",
    "liability": "Current Liabilities", "non-current liability": "Non-current Liabilities", "equity": "Equity",
    "revenue": "Trading Income", "sales": "Trading Income", "other income": "Other Income",
    "direct costs": "Cost of Sales", "expense": "Operating Expenses", "overhead": "Operating Expenses",
    "overheads": "Operating Expenses", "depreciation": "Operating Expenses",
}


class ImportError_(LedgerError):
    pass


# ---------------------------------------------------------------------------------------- #
# reading files
# ---------------------------------------------------------------------------------------- #
def read_table(path: str | Path, sheet: str | None = None) -> list[list]:
    """Rows of a CSV (any common delimiter) or the first / named sheet of an Excel file."""
    p = Path(path)
    if p.suffix.lower() in (".xlsx", ".xlsm", ".xltx", ".xltm"):
        import openpyxl
        wb = openpyxl.load_workbook(p, data_only=True, read_only=True)
        ws = wb[sheet] if sheet and sheet in wb.sheetnames else wb[wb.sheetnames[0]]
        return [list(r) for r in ws.iter_rows(values_only=True)]
    raw = p.read_bytes()
    for enc in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            text = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    return [row for row in csv.reader(io.StringIO(text), dialect)]


def _norm(h) -> str:
    return " ".join(str(h or "").replace("*", "").replace("_", " ").strip().lower().split())


def _cell(row: list, i: int | None):
    if i is None or i >= len(row):
        return None
    return row[i]


# ---------------------------------------------------------------------------------------- #
# balances
# ---------------------------------------------------------------------------------------- #
HEAD = {
    "code": {"code", "account code", "acc code", "a/c code", "account no", "account number", "acct code", "gl code"},
    "name": {"name", "account name", "account", "description", "particulars", "account code - name", "account title"},
    "report_code": {"report code", "reporting code", "report code name"},
    "type": {"type", "account type", "class"},
    "debit": {"debit", "dr", "debits", "debit amount"},
    "credit": {"credit", "cr", "credits", "credit amount"},
    "ytd_debit": {"ytd debit", "debit - ytd", "debit ytd", "year to date debit", "ytd debits"},
    "ytd_credit": {"ytd credit", "credit - ytd", "credit ytd", "year to date credit", "ytd credits"},
    "balance": {"balance", "amount", "closing balance", "net", "value", "cy", "current year", "this year", "net movement"},
    "py": {"py", "prior year", "comparative", "last year", "previous year"},
}


@dataclass
class BalanceRow:
    line: int
    code: str
    name: str
    amount: Decimal                 # debit positive
    py: Decimal | None = None       # debit positive
    report_code: str = ""
    acct_type: str = ""
    section: str = ""
    note: str = ""


def detect_balance_columns(rows: list[list]) -> tuple[int, dict]:
    """Find the header row and the column of each field."""
    best = (-1, {})
    for i, row in enumerate(rows[:30]):
        cols: dict = {}
        years = []
        for j, h in enumerate(row):
            n = _norm(h)
            if not n:
                continue
            for key, names in HEAD.items():
                if n in names and key not in cols:
                    cols[key] = j
                    break
            else:
                if build_pack._is_period_header(h):
                    years.append(j)
        if years and "balance" not in cols:
            cols["balance"] = years[0]
            if len(years) > 1 and "py" not in cols:
                cols["py"] = years[1]
        has_amount = any(k in cols for k in ("debit", "credit", "ytd_debit", "ytd_credit", "balance"))
        has_label = "code" in cols or "name" in cols
        if has_amount and has_label and len(cols) > len(best[1]):
            best = (i, cols)
    if best[0] < 0:
        raise ImportError_("could not find a header row with an account column and an amount column "
                           "(e.g. Code, Name, Debit, Credit - or Account, Balance)")
    return best


NAME_CODE_RE = re.compile(r"^(?P<name>.*\S)\s*\((?P<code>[A-Za-z0-9][A-Za-z0-9.\-]{0,9})\)\s*$")


def split_code_name(code_cell, name_cell) -> tuple[str, str]:
    code = str(code_cell).strip() if code_cell not in (None, "") else ""
    name = str(name_cell).strip() if name_cell not in (None, "") else ""
    if isinstance(code_cell, float) and code_cell.is_integer():
        code = str(int(code_cell))
    if not code and name:
        m = NAME_CODE_RE.match(name)
        if m:
            return m.group("code"), m.group("name")
        c, n = fr_engine.strip_code(name)
        if c:
            return c, n
    if code and not name:
        c, n = fr_engine.strip_code(code)
        if c:
            return c, n
    return code, name


def parse_balances(rows: list[list], sign: str = "auto") -> tuple[list[BalanceRow], dict]:
    """
    Rows -> BalanceRows (amounts debit positive).  sign: 'auto' | 'debit' (positive = debit) |
    'natural' (positive = the account's normal side: income, liabilities and equity are credits).
    """
    hi, cols = detect_balance_columns(rows)
    out: list[BalanceRow] = []
    section = ""
    single = "balance" in cols and not any(k in cols for k in ("debit", "credit", "ytd_debit", "ytd_credit"))
    for idx, row in enumerate(rows[hi + 1:], start=hi + 2):
        row = list(row)
        code, name = split_code_name(_cell(row, cols.get("code")), _cell(row, cols.get("name")))
        if not code and not name:
            continue
        label = name or code
        key = label.lower().rstrip(":").strip()
        if "ytd_debit" in cols or "ytd_credit" in cols:
            dr, cr = _cell(row, cols.get("ytd_debit")), _cell(row, cols.get("ytd_credit"))
        else:
            dr, cr = _cell(row, cols.get("debit")), _cell(row, cols.get("credit"))
        bal, pyv = _cell(row, cols.get("balance")), _cell(row, cols.get("py"))
        amounts = [v for v in (dr, cr, bal, pyv) if v not in (None, "") and is_amount(v)]
        if not amounts:
            if key in chartlib.SECTIONS:
                section = label
            continue
        if not code and (build_pack.SKIP_LABELS.match(label) or key.startswith("total")):
            continue
        if single:
            amount = to_dec(bal)
        else:
            amount = to_dec(dr) - to_dec(cr)
        py = to_dec(pyv) if pyv not in (None, "") and is_amount(pyv) else None
        rc = str(_cell(row, cols.get("report_code")) or "").strip().upper()
        if rc and slot_of(rc) is None:
            rc = ""
        out.append(BalanceRow(line=idx, code=code, name=name or code, amount=amount, py=py, report_code=rc,
                              acct_type=str(_cell(row, cols.get("type")) or "").strip(), section=section))
    info = {"header_row": hi + 1, "columns": cols, "single_column": single, "sign": "debit"}
    if single:
        info["sign"] = sign
    return out, info


def _row_nature(r: BalanceRow, ledger: Ledger) -> str:
    a = ledger.chart.get(r.code) if r.code else None
    rc = r.report_code or (a.report_code if a else "") or default_report_code(r.acct_type)
    if not rc and r.section:
        side, sec, nat = chartlib.section_info(r.section)
        if nat:
            return nat
    return nature(rc, r.name) if rc else "dr"


def apply_sign(rows: list[BalanceRow], info: dict, ledger: Ledger) -> None:
    """Single-column files: decide whether positive means debit or 'natural side', then convert."""
    if not info.get("single_column"):
        return
    mode = info.get("sign", "auto")
    if mode == "auto":
        s_debit = sum((r.amount for r in rows), ZERO)
        s_nat = sum(((r.amount if _row_nature(r, ledger) == "dr" else -r.amount) for r in rows), ZERO)
        mode = "natural" if abs(s_nat) < abs(s_debit) else "debit"
        info["sign_detected"] = mode
    info["sign"] = mode
    if mode == "natural":
        for r in rows:
            if _row_nature(r, ledger) == "cr":
                r.amount = -r.amount
                if r.py is not None:
                    r.py = -r.py


# ---------------------------------------------------------------------------------------- #
# planning: which account each line posts to
# ---------------------------------------------------------------------------------------- #
@dataclass
class PlanLine:
    line: int
    source_code: str
    source_name: str
    amount: str                    # debit positive, string for JSON
    py: str | None = None
    action: str = "post"           # post | new | skip
    account: str = ""              # target account code
    account_name: str = ""
    report_code: str = ""
    prints_as: str = ""
    confidence: str = ""
    reason: str = ""
    review: bool = False


@dataclass
class Plan:
    kind: str                      # balances | workpaper
    fy: int
    mode: str = "year_end"         # year_end | opening
    source: str = ""
    lines: list[PlanLine] = field(default_factory=list)
    info: dict = field(default_factory=dict)
    warnings: list[str] = field(default_factory=list)
    complete: bool = True          # accounts not in the file are set to nil
    include_py: bool = False

    def to_dict(self) -> dict:
        return asdict(self)

    @classmethod
    def from_dict(cls, d: dict) -> "Plan":
        lines = [PlanLine(**x) for x in d.get("lines", [])]
        return cls(**{**{k: v for k, v in d.items() if k != "lines"}, "lines": lines})

    @property
    def total(self) -> Decimal:
        return sum((to_dec(pl.amount) for pl in self.lines if pl.action != "skip"), ZERO)

    @property
    def total_py(self) -> Decimal:
        return sum((to_dec(pl.py) for pl in self.lines if pl.action != "skip" and pl.py is not None), ZERO)


CYE_RE = re.compile(r"current year (earnings|profit)|profit for the (year|period)|net income for the year", re.I)


def plan_balances(ledger: Ledger, rows: list[BalanceRow], fy: int, mode: str = "year_end", source: str = "",
                  info: dict | None = None) -> Plan:
    plan = Plan(kind="balances", fy=fy, mode=mode, source=source, info=info or {})
    matcher = ledger.chart.matcher()
    for r in rows:
        pl = PlanLine(line=r.line, source_code=r.code, source_name=r.name, amount=as_str(r.amount),
                      py=as_str(r.py) if r.py is not None else None)
        a = ledger.chart.get(r.code) if r.code else None
        rc_in = r.report_code
        if CYE_RE.search(r.name or "") or rc_in.startswith("EQU.RET.CUR") or (a and a.report_code.startswith("EQU.RET.CUR")):
            pl.action, pl.reason = "skip", "current year earnings - the ledger works it out from the P&L"
        elif a is not None and not a.system:
            pl.account, pl.account_name, pl.report_code = a.code, a.name, a.report_code
            pl.confidence, pl.reason = "code", "same account code in the client's chart"
            if rc_in and rc_in != a.report_code:
                pl.reason += f" (file says {rc_in}; the chart's {a.report_code} is used - change it in Accounts)"
        else:
            section = r.section or XERO_TYPE_SECTION.get((r.acct_type or "").lower(), "")
            if not section and rc_in:
                section = _section_for_code(rc_in)
            exact = ledger.chart.by_name(r.name)
            # a row without a code is known only by its name; a row with its own code and report code is
            # the file's own account (it becomes a new account rather than joining a same-named one)
            if exact is not None and not exact.system and (not rc_in or not r.code):
                pl.account, pl.account_name, pl.report_code = exact.code, exact.name, exact.report_code
                pl.confidence, pl.reason = "name", "same account name in the client's chart"
                if rc_in and rc_in != exact.report_code:
                    pl.reason += f" (file says {rc_in}; the chart's {exact.report_code} is used)"
            else:
                if rc_in:
                    rc, conf, why, review = rc_in, "file", "report code from the file", False
                else:
                    side = "pl" if is_pl_account("", r.acct_type) else None
                    m = matcher.match(r.name, section or None, r.code or None, side_hint=side)
                    rc, conf, why, review = m.report_code, m.confidence, m.reason, m.review
                    if not section:
                        # no section header, account type or report code in the file: check these by eye
                        review = True
                        why += " (no account type in the file - check)"
                    if m.account is not None and m.confidence in ("chart-exact", "chart-code") and not r.code:
                        pl.account, pl.account_name, pl.report_code = m.account.code, m.account.name, m.account.report_code
                        pl.confidence, pl.reason = m.confidence, m.reason
                if not pl.account:
                    pl.action = "new"
                    pl.account = r.code if (r.code and not ledger.chart.get(r.code) and
                                            re.fullmatch(r"[A-Za-z0-9][A-Za-z0-9.\-]{0,9}", r.code)) else ""
                    pl.account_name = r.name
                    pl.report_code = rc
                    pl.confidence, pl.reason, pl.review = conf, f"new account - {why}", review
        if pl.report_code:
            pl.prints_as = describe(pl.report_code, pl.account_name)
        plan.lines.append(pl)
    _allocate_codes(ledger, plan)
    _plan_checks(ledger, plan)
    return plan


def _section_for_code(rc: str) -> str:
    rc = rc.upper()
    for pref, sec in (("ASS.CUR.CAS", "Bank"), ("ASS.CUR", "Current Assets"), ("ASS.NCA.FIX", "Fixed Assets"),
                      ("ASS", "Non-current Assets"), ("LIA.CUR", "Current Liabilities"), ("LIA", "Non-current Liabilities"),
                      ("EQU", "Equity"), ("REV.TRA", "Trading Income"), ("REV", "Other Income"),
                      ("EXP.COS", "Cost of Sales"), ("EXP", "Operating Expenses")):
        if rc.startswith(pref):
            return sec
    return ""


def _allocate_codes(ledger: Ledger, plan: Plan) -> None:
    """New accounts without a usable code get one next to the closest chart account."""
    taken = {a.code for a in ledger.chart.accounts}
    for pl in plan.lines:
        if pl.action != "new":
            continue
        if pl.account and pl.account not in taken:
            taken.add(pl.account)
            continue
        base = _nearest_code(ledger, pl.report_code) or "9000"
        i = 1
        while True:
            cand = (f"{base}.{i}" if base[-1].isdigit() else f"{base}{i}")[:10]
            if cand not in taken:
                break
            i += 1
        pl.account = cand
        taken.add(cand)


def _nearest_code(ledger: Ledger, rc: str) -> str:
    rc = (rc or "").upper()
    parts = rc.split(".")
    for i in range(len(parts), 0, -1):
        pref = ".".join(parts[:i])
        cands = [a for a in ledger.chart.usable() if a.report_code == pref or a.report_code.startswith(pref + ".")]
        if cands:
            return sorted(cands, key=lambda a: code_sort_key(a.code))[0].code
    return ""


def _plan_checks(ledger: Ledger, plan: Plan) -> None:
    diff = plan.total
    if diff and abs(diff) <= ROUNDING_TOLERANCE:
        plan.warnings.append(f"the file is out of balance by {diff:,.2f} (rounding) - the difference will be "
                             f"posted to the rounding account")
    elif diff:
        plan.warnings.append(f"the file does not balance: debits exceed credits by {diff:,.2f}. Fix the file "
                             f"(a missing line?) or tick 'post the difference to suspense'.")
    if any(pl.py is not None for pl in plan.lines):
        dpy = plan.total_py
        if dpy and abs(dpy) > ROUNDING_TOLERANCE:
            plan.warnings.append(f"the prior-year column does not balance (difference {dpy:,.2f})")
    if plan.mode == "opening":
        pls = [pl for pl in plan.lines if pl.action != "skip" and is_pl_account(pl.report_code)]
        if pls:
            plan.warnings.append(f"{len(pls)} profit and loss line(s) in an opening-balance import - opening balances are "
                                 f"balance-sheet only; post last year's result to retained earnings instead")
    if ledger.is_locked(plan.fy):
        plan.warnings.append(f"FY{plan.fy} is locked - unlock it before importing")


def recheck(ledger: Ledger, plan: Plan) -> list[str]:
    """
    After the user has changed lines in the preview (post to another account, new account, skip):
    describe each line again, give new accounts a code, and redo the checks.  Returns the problems
    that stop the import.
    """
    problems = []
    new_codes: dict[str, str] = {}
    for pl in plan.lines:
        pl.account = (pl.account or "").strip()
        pl.report_code = (pl.report_code or "").strip().upper()
        where = f"line {pl.line} ({pl.source_name})"
        if pl.action == "post":
            a = ledger.chart.get(pl.account)
            if a is None:
                problems.append(f"{where}: account {pl.account or '(blank)'} is not in the chart - pick one or make it a new account")
            elif a.system:
                problems.append(f"{where}: {a.code} {a.name} is a Xero system account")
            else:
                pl.account_name, pl.report_code = a.name, a.report_code
        elif pl.action == "new":
            if pl.account and ledger.chart.get(pl.account):
                problems.append(f"{where}: account {pl.account} already exists - post to it or use another code")
            if pl.account and pl.account in new_codes and new_codes[pl.account] != pl.account_name:
                problems.append(f"{where}: new account {pl.account} is also used for '{new_codes[pl.account]}'")
            if pl.account:
                new_codes[pl.account] = pl.account_name
            if not (pl.account_name or "").strip():
                pl.account_name = pl.source_name
            if slot_of(pl.report_code) is None:
                problems.append(f"{where}: report code {pl.report_code or '(blank)'} is not a known report code")
        if pl.report_code and pl.action != "skip":
            pl.prints_as = describe(pl.report_code, pl.account_name)
    _allocate_codes(ledger, plan)
    plan.warnings = []
    _plan_checks(ledger, plan)
    return problems


# ---------------------------------------------------------------------------------------- #
# applying
# ---------------------------------------------------------------------------------------- #
def apply_plan(ledger: Ledger, plan: Plan, allow_suspense: bool = False, replace_imports: bool = True) -> dict:
    """Create the new accounts and post the import journal(s).  Returns what was done."""
    ledger.store.require_lock(ledger.slug)
    blocking = [w for w in plan.warnings if "is locked" in w]
    if blocking:
        raise ValidationError(blocking)
    diff = plan.total
    if diff and abs(diff) > ROUNDING_TOLERANCE and not allow_suspense:
        raise ValidationError([f"the file does not balance (difference {diff:,.2f})"])
    created = []
    for pl in plan.lines:
        if pl.action == "new" and not ledger.chart.get(pl.account):
            try:
                ledger.chart.add(pl.account, pl.account_name, pl.report_code)
            except ChartError as e:
                raise ValidationError([f"line {pl.line} ({pl.source_name}): {e}"]) from None
            ledger.store.audit(ledger.slug, "account.added", {"code": pl.account, "name": pl.account_name,
                                                              "report_code": pl.report_code, "why": f"import {plan.source}"})
            created.append(pl.account)
    done = {"created_accounts": created, "journals": []}
    targets_cy = _targets(plan, "amount")
    if plan.include_py and any(pl.py is not None for pl in plan.lines):
        targets_py = _targets(plan, "py")
        j = _post_targets(ledger, plan.fy - 1, targets_py, plan, "year_end", allow_suspense, replace_imports,
                          f"Balances imported from {plan.source} (prior year)")
        if j:
            done["journals"].append(j.number)
    j = _post_targets(ledger, plan.fy, targets_cy, plan, plan.mode, allow_suspense, replace_imports,
                      f"{'Opening balances' if plan.mode == 'opening' else 'Balances'} imported from {plan.source}")
    if j:
        done["journals"].append(j.number)
    return done


def _targets(plan: Plan, attr: str) -> dict[str, Decimal]:
    out: dict[str, Decimal] = {}
    for pl in plan.lines:
        if pl.action == "skip":
            continue
        v = getattr(pl, attr)
        if v is None:
            continue
        out[pl.account] = out.get(pl.account, ZERO) + to_dec(v)
    return out


def _post_targets(ledger: Ledger, fy: int, targets: dict[str, Decimal], plan: Plan, mode: str, allow_suspense: bool,
                  replace_imports: bool, narration: str) -> Journal | None:
    ledger.ensure_year(fy)
    if ledger.is_locked(fy):
        raise ValidationError([f"FY{fy} is locked - unlock it before importing"])
    start, end = ledger.fy_range(fy)
    on = start if mode == "opening" else end
    kind = "opening" if mode == "opening" else "import"
    if replace_imports:
        for old in ledger.journals(fy):
            if old.kind == kind and old.source.get("import"):
                ledger.void_journal(fy, old.number, f"replaced by a new import of {plan.source}")
    # what the ledger holds now, at the import date (opening carried forward + other journals)
    if mode == "opening":
        current = ledger.opening(fy)
        for k, v in ledger.movements(fy, upto=start).items():
            current[k] = current.get(k, ZERO) + v
    else:
        current = ledger.closing(fy)
    accounts = set(targets)
    if plan.complete:
        if mode == "opening":
            accounts |= {k for k in current if not ledger.is_pl(k)}
        else:
            accounts |= set(current)
    lines = []
    for code in sorted(accounts, key=code_sort_key):
        delta = targets.get(code, ZERO) - current.get(code, ZERO)
        if delta:
            lines.append(Line(code, debit=delta if delta > 0 else ZERO, credit=-delta if delta < 0 else ZERO,
                              description="imported balance" if code in targets else "not in the file - set to nil"))
    if not lines:
        return None
    diff = sum((ln.amount for ln in lines), ZERO)
    if diff:
        if abs(diff) <= ROUNDING_TOLERANCE:
            acct, why = ledger.rounding_code(), "rounding in the imported file"
        elif allow_suspense:
            acct, why = _suspense_code(ledger), "out of balance in the imported file - clear this"
        else:
            raise ValidationError([f"FY{fy}: the import does not balance (difference {diff:,.2f})"])
        lines.append(Line(acct, debit=-diff if diff < 0 else ZERO, credit=diff if diff > 0 else ZERO, description=why))
    j = Journal(date=on, narration=narration, kind=kind, lines=lines,
                source={"import": plan.kind, "file": plan.source, "mode": mode})
    return ledger.save_journal(j, fy, post=True, action="import.posted")


def _suspense_code(ledger: Ledger) -> str:
    cands = [a for a in ledger.chart.usable() if re.search(r"suspense", a.name, re.I)]
    if cands:
        return cands[0].code
    code = "10100" if not ledger.chart.get("10100") else ledger.chart.next_code("10100")
    ledger.chart.add(code, "Suspense", "ASS.CUR", "Current Asset")
    ledger.store.audit(ledger.slug, "account.added", {"code": code, "name": "Suspense", "report_code": "ASS.CUR"})
    return code


# ---------------------------------------------------------------------------------------- #
# journals CSV
# ---------------------------------------------------------------------------------------- #
JHEAD = {
    "narration": {"narration", "memo", "journal narration", "journal description"},
    "date": {"date", "journal date", "transaction date"},
    "reference": {"reference", "journal", "journal no", "journal number", "ref", "number", "journal ref"},
    "account": {"accountcode", "account code", "account", "code", "gl code"},
    "description": {"description", "line description", "details"},
    "amount": {"amount", "value", "net amount"},
    "debit": {"debit", "dr"},
    "credit": {"credit", "cr"},
}


def parse_journal_rows(rows: list[list]) -> tuple[list[Journal], list[str]]:
    """Group a journal CSV into journals.  Xero format: Amount > 0 is a debit, < 0 a credit."""
    hi, cols = None, {}
    squash = {key: {x.replace(" ", "") for x in names} for key, names in JHEAD.items()}
    for i, row in enumerate(rows[:20]):
        c = {}
        for j, h in enumerate(row):
            n = _norm(h).replace(" ", "")
            for key in JHEAD:
                if n and n in squash[key] and key not in c:
                    c[key] = j
                    break
        if "account" in c and "date" in c and ("amount" in c or "debit" in c or "credit" in c):
            hi, cols = i, c
            break
    if hi is None:
        raise ImportError_("could not find the journal columns (need Date, Account Code and Amount or Debit / Credit)")
    groups: dict[tuple, Journal] = {}
    order = []
    errors = []
    for idx, row in enumerate(rows[hi + 1:], start=hi + 2):
        row = list(row)
        acct = str(_cell(row, cols.get("account")) or "").strip()
        if isinstance(_cell(row, cols.get("account")), float) and float(_cell(row, cols.get("account"))).is_integer():
            acct = str(int(_cell(row, cols.get("account"))))
        if not acct:
            continue
        try:
            d = parse_date(_cell(row, cols.get("date")))
        except ValueError as e:
            errors.append(f"row {idx}: {e}")
            continue
        narr = str(_cell(row, cols.get("narration")) or "").strip()
        ref = str(_cell(row, cols.get("reference")) or "").strip()
        if "amount" in cols:
            amt = to_dec(_cell(row, cols["amount"]))
            dr, cr = (amt, ZERO) if amt > 0 else (ZERO, -amt)
        else:
            dr, cr = to_dec(_cell(row, cols.get("debit"))), to_dec(_cell(row, cols.get("credit")))
        key = (ref,) if ref else (narr, d)
        if key not in groups:
            groups[key] = Journal(date=d, narration=narr or ref, reference=ref, lines=[], kind="manual",
                                  source={"import": "journals"})
            order.append(key)
        groups[key].lines.append(Line(acct, debit=dr, credit=cr,
                                      description=str(_cell(row, cols.get("description")) or "").strip()))
    return [groups[k] for k in order], errors


def import_journals(ledger: Ledger, journals: list[Journal], source: str = "", post: bool = True) -> list[str]:
    """Validate every journal first; post all or none."""
    problems = []
    for j in journals:
        fy = ledger.client.fy_of(j.date)
        ledger.ensure_year(fy)
        errs = ledger.validate(j, fy)
        if errs:
            problems.append(f"{j.narration or j.reference} ({j.date:%d/%m/%Y}): " + "; ".join(errs))
    if problems and post:
        raise ValidationError(problems)
    numbers = []
    for j in journals:
        j.source = {**j.source, "file": source}
        fy = ledger.client.fy_of(j.date)
        numbers.append(ledger.save_journal(j, fy, post=post, action="import.journal").number)
    return numbers


# ---------------------------------------------------------------------------------------- #
# Fortis workpaper
# ---------------------------------------------------------------------------------------- #
def plan_workpaper(ledger: Ledger, path: str | Path, fy: int, pl_sheet: str = "Profit & Loss,PL",
                   bs_sheet: str = "Balance Sheet,BS") -> Plan:
    """Read a Fortis workpaper with the skill's reader and classify each line by nature (report code)."""
    import openpyxl
    names = openpyxl.load_workbook(path, read_only=True).sheetnames
    pl_names = [x.strip() for x in pl_sheet.split(",")]
    stop_at = None
    if not any(n in names for n in pl_names) and "Business" in names:
        pl_rows, bs_rows = build_pack.split_side_by_side(build_pack.read_rows_xlsx(str(path), ["Business"]))
        stop_at = r"^net (profit|loss)\b"
    else:
        pl_rows = build_pack.read_rows_xlsx(str(path), pl_names)
        bs_rows = build_pack.read_rows_xlsx(str(path), [x.strip() for x in bs_sheet.split(",")])
    pl = build_pack.parse_statement(pl_rows, "pl", stop_at)
    bs = build_pack.parse_statement(bs_rows, "bs")
    accts = build_pack.post_lines(pl + bs, ledger.chart.chart, REPORT_MAP, {}, "auto", ledger.client.type)
    plan = Plan(kind="workpaper", fy=fy, mode="year_end", source=Path(path).name, include_py=False)
    for a in accts:
        rc = a["report_code"]
        sign = 1 if nature(rc, a["name"]) == "dr" else -1
        amount = to_dec(a["cy"]) * sign
        py = to_dec(a["py"]) * sign
        m = a.get("mapping", {})
        pl_ = PlanLine(line=int(re.sub(r"\D", "", a["source"]) or 0), source_code=a.get("code") or "",
                       source_name=a["name"], amount=as_str(amount), py=as_str(py), report_code=rc,
                       confidence=m.get("confidence", ""), reason=m.get("reason", ""), review=bool(m.get("review")))
        if rc.startswith("EQU.RET.CUR"):
            pl_.action, pl_.reason = "skip", "current year earnings - the ledger works it out from the P&L"
        else:
            existing = ledger.chart.get(a.get("chart_code") or "")
            same_name = existing is not None and " ".join(existing.name.lower().split()) == " ".join(a["name"].lower().split())
            if existing is not None and (same_name or m.get("confidence") == "chart-code"):
                pl_.account, pl_.account_name, pl_.report_code = existing.code, existing.name, existing.report_code
            else:
                pl_.action = "new"
                wp_code = (a.get("code") or "").strip()
                pl_.account = wp_code if wp_code and not ledger.chart.get(wp_code) else ""
                pl_.account_name = a["name"]
        pl_.prints_as = describe(pl_.report_code, pl_.account_name or a["name"])
        plan.lines.append(pl_)
    plan.include_py = any(to_dec(x.py) for x in plan.lines if x.py is not None) and not ledger.journals(fy - 1)
    _allocate_codes(ledger, plan)
    _plan_checks(ledger, plan)
    return plan


AS_AT_RE = re.compile(r"\b(?:as at|year ended|period ended|year ending)\s+(\d{1,2}\s+[A-Za-z]+\s+\d{4}|"
                      r"\d{1,2}/\d{1,2}/\d{4}|\d{4}-\d{2}-\d{2})", re.I)


def detect_as_at(rows: list[list]):
    """The 'As at 30 June 2026' date in the title rows of a Xero report export, or None."""
    import datetime as _dt
    for row in rows[:10]:
        for v in row:
            if isinstance(v, _dt.datetime):
                continue
            m = AS_AT_RE.search(str(v or ""))
            if m:
                try:
                    return parse_date(m.group(1))
                except ValueError:
                    continue
    return None


def plan_from_file(ledger: Ledger, path: str | Path, fy: int, mode: str = "year_end", sign: str = "auto",
                   sheet: str | None = None) -> Plan:
    rows = read_table(path, sheet)
    parsed, info = parse_balances(rows, sign)
    apply_sign(parsed, info, ledger)
    plan = plan_balances(ledger, parsed, fy, mode, Path(path).name, info)
    plan.include_py = any(r.py is not None for r in parsed) and not ledger.journals(fy - 1)
    return plan
